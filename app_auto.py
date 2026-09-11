import json
import os
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

import streamlit as st
from crewai import Agent, Task, Crew, Process, LLM

AGENTS_FILE = Path(__file__).with_name("agents.json")
HISTORY_FILE = Path(__file__).with_name("workflow_history.json")
DEFAULT_AGENTS = [
    {"role": "Research Analyst", "goal": "Find reliable information and turn it into clear, useful insights", "backstory": "You are a careful researcher who compares sources, spots patterns, and explains complex topics plainly."},
    {"role": "Content Strategist", "goal": "Shape information into focused content for a specific audience", "backstory": "You are an experienced editor who understands audience needs, structure, tone, and persuasive communication."},
    {"role": "Data Analyst", "goal": "Interpret data, identify meaningful trends, and communicate practical conclusions", "backstory": "You are a methodical analyst who checks assumptions, looks for outliers, and presents evidence-driven recommendations."},
    {"role": "Quality Reviewer", "goal": "Check work for accuracy, completeness, clarity, and actionable improvements", "backstory": "You are a constructive reviewer who catches gaps and inconsistencies while keeping the final result easy to use."},
]
WEB_SEARCH_TERMS = (
    "web", "website", "websites", "internet", "online", "browse", "browser",
    "crawl", "crawling", "scrape", "scraping", "search", "google", "linkedin",
    "competitor", "competitors", "current", "latest", "recent", "news", "market",
    "source", "sources", "research", "job listing", "job listings", "live data",
)
LLM_PROVIDERS = {
    "Google Gemini": {
        "env_key": "GEMINI_API_KEY",
        "models": ["gemini/gemini-3.6-flash"],
        "requires_key": True,
    },
    "OpenAI": {
        "env_key": "OPENAI_API_KEY",
        "models": ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini", "gpt-4.1"],
        "requires_key": True,
    },
    "Anthropic": {
        "env_key": "ANTHROPIC_API_KEY",
        "models": ["anthropic/claude-haiku-4-5", "anthropic/claude-sonnet-4-5"],
        "requires_key": True,
    },
    "Groq": {
        "env_key": "GROQ_API_KEY",
        "models": ["groq/llama-3.3-70b-versatile", "groq/llama-3.1-8b-instant"],
        "requires_key": True,
    },
    "OpenRouter": {
        "env_key": "OPENROUTER_API_KEY",
        "models": ["openrouter/openai/gpt-4o-mini", "openrouter/google/gemini-3.6-flash"],
        "requires_key": True,
    },
    "Mistral": {
        "env_key": "MISTRAL_API_KEY",
        "models": ["mistral/mistral-small-latest", "mistral/mistral-large-latest"],
        "requires_key": True,
    },
    "Hugging Face": {
        "env_key": "HF_TOKEN",
        "models": ["huggingface/meta-llama/Llama-3.1-8B-Instruct", "huggingface/mistralai/Mistral-7B-Instruct-v0.3"],
        "requires_key": True,
    },
    "Ollama (local)": {
        "env_key": "",
        "models": ["ollama/llama3.2", "ollama/qwen2.5:7b", "ollama/mistral"],
        "requires_key": False,
    },
}


def workflow_requires_web_search(text):
    normalized_text = (text or "").lower()
    return any(term in normalized_text for term in WEB_SEARCH_TERMS)


def create_selected_llm(provider_name, model_name, api_key):
    provider = LLM_PROVIDERS[provider_name]
    if provider["requires_key"] and not api_key.strip():
        return None

    options = {"model": model_name}
    if provider["requires_key"]:
        options["api_key"] = api_key.strip()
    return LLM(**options)


def load_saved_agents():
    if not AGENTS_FILE.exists():
        return [dict(agent) for agent in DEFAULT_AGENTS]
    try:
        with AGENTS_FILE.open("r", encoding="utf-8") as file:
            saved_agents = json.load(file)
        return saved_agents if isinstance(saved_agents, list) else []
    except (OSError, json.JSONDecodeError):
        return [dict(agent) for agent in DEFAULT_AGENTS]


def save_agents(agents):
    with AGENTS_FILE.open("w", encoding="utf-8") as file:
        json.dump(agents, file, indent=2)


def load_history():
    if not HISTORY_FILE.exists():
        return []
    try:
        with HISTORY_FILE.open("r", encoding="utf-8") as file:
            history = json.load(file)
        return history if isinstance(history, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def save_history(history):
    with HISTORY_FILE.open("w", encoding="utf-8") as file:
        json.dump(history, file, indent=2)


def add_history_entry(workflow_brief):
    clean_text = (workflow_brief or "").strip()
    if not clean_text:
        return

    history = list(st.session_state.get("history", []))
    history = [entry for entry in history if (entry.get("prompt") or "").strip() != clean_text]
    history.insert(0, {
        "prompt": clean_text,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
    st.session_state.history = history[:20]
    save_history(st.session_state.history)


def extract_document_text(uploaded_file):
    """Extract text from the supported document formats."""
    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith((".txt", ".md", ".csv", ".json")):
        return file_bytes.decode("utf-8", errors="replace")

    if file_name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if file_name.endswith(".docx"):
        from docx import Document

        document = Document(BytesIO(file_bytes))
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs).strip()

    raise ValueError("Unsupported file type. Use TXT, MD, CSV, JSON, PDF, or DOCX.")


def create_docx_download(text):
    from docx import Document

    document = Document()
    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def parse_json_from_text(raw_text):
    if raw_text is None:
        raise ValueError("No response from the model.")

    if isinstance(raw_text, dict):
        return raw_text

    text = str(raw_text).strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text, flags=re.IGNORECASE)

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            raise ValueError("Model did not return valid JSON.")
        data = json.loads(match.group(0))

    return data


def normalize_workflow_payload(payload):
    if isinstance(payload, dict):
        agents = payload.get("agents") or payload.get("roles") or []
        tasks = payload.get("tasks") or payload.get("workflow") or []
    elif isinstance(payload, list):
        agents = []
        tasks = payload
    else:
        raise ValueError("Workflow payload must be a JSON object with agents and tasks.")

    cleaned_agents = []
    for item in agents:
        if not isinstance(item, dict):
            continue
        role = (item.get("role") or item.get("name") or "").strip()
        goal = (item.get("goal") or item.get("objective") or "").strip()
        backstory = (item.get("backstory") or item.get("style") or "").strip()
        if role and goal and backstory:
            cleaned_agents.append({"role": role, "goal": goal, "backstory": backstory})

    cleaned_tasks = []
    for item in tasks:
        if not isinstance(item, dict):
            continue
        description = (item.get("description") or item.get("task") or "").strip()
        expected_output = (item.get("expected_output") or item.get("deliverable") or "").strip()
        agent_role = (item.get("agent_role") or item.get("assigned_role") or item.get("agent") or "").strip()

        if description and expected_output:
            cleaned_tasks.append({
                "description": description,
                "expected_output": expected_output,
                "agent_role": agent_role or (cleaned_agents[0]["role"] if cleaned_agents else "Workflow Specialist"),
            })

    if not cleaned_agents and cleaned_tasks:
        unique_roles = []
        seen = set()
        for task in cleaned_tasks:
            role = task["agent_role"]
            if role and role not in seen:
                unique_roles.append({
                    "role": role,
                    "goal": f"Complete the task: {task['description'][:80]}",
                    "backstory": "You are a specialist agent created specifically for this workflow."
                })
                seen.add(role)
        cleaned_agents = unique_roles

    if cleaned_agents and not cleaned_tasks:
        cleaned_tasks = [{
            "description": f"Execute the workflow for: {cleaned_agents[0]['goal']}",
            "expected_output": "A concise, actionable outcome that fulfills the requested workflow.",
            "agent_role": cleaned_agents[0]["role"]
        }]

    return cleaned_agents, cleaned_tasks


def generate_workflow_from_brief(workflow_brief, llm):
    if not llm:
        raise ValueError("Add the provider API key before generating a workflow.")

    prompt = f"""
You are a workflow architect. Create a realistic multi-agent work plan from this brief.

USER BRIEF:
{workflow_brief}

Design the workflow entirely from the user's brief. Do not assume this is a resume,
career, LinkedIn, research, software, marketing, or any other specific type of task
unless the brief explicitly says so. The agent roles, goals, tasks, and deliverables
must be directly relevant to the user's requested outcome.

Return a valid JSON object with exactly these keys:
{{
  "agents": [
        {{"role": "A role chosen from the user's brief", "goal": "A goal directly related to the requested outcome", "backstory": "A working style appropriate for this role"}}
  ],
  "tasks": [
        {{"description": "A specific task derived from the user's brief", "expected_output": "The concrete deliverable from this task", "agent_role": "One of the generated agent roles"}}
  ]
}}

Rules:
- Create the smallest useful team: usually 2 to 5 agents and 2 to 5 tasks.
- Choose role names based on the domain, actions, and deliverables in the brief.
- Break the brief into a logical order of tasks with clear dependencies.
- Make every task contribute directly to the requested outcome.
- Use the attached document as source material when one is provided at execution time.
- Match each task to an existing agent role.
- Keep the response valid JSON only; no markdown fences, no narration, no comments.
- Make the workflow realistic, actionable, and specific to this brief.
"""
    response = llm.call(prompt)
    payload = parse_json_from_text(response)
    return normalize_workflow_payload(payload)


st.set_page_config(page_title="Auto Workflow Builder", layout="wide")
st.title("Workflow Studio Auto Builder")
st.subheader("Describe the workflow once; let the system propose the roles and tasks")

with st.sidebar:
    st.header("Model settings")
    provider_name = st.selectbox("LLM provider", options=list(LLM_PROVIDERS))
    provider_config = LLM_PROVIDERS[provider_name]
    selected_model = st.selectbox("Select model", options=provider_config["models"])

    if provider_config["requires_key"]:
        api_key_input = st.text_input(
            f"{provider_name} API key",
            type="password",
            value=os.getenv(provider_config["env_key"], ""),
            help="Used only in this Streamlit session unless supplied by the environment.",
        )
    else:
        api_key_input = ""
        st.caption("Requires the Ollama app running locally with the selected model pulled.")

    process_type = st.radio(
        "Workflow mode",
        options=["Sequential", "Hierarchical"],
        index=0,
        help="Choose whether tasks run in order or are coordinated by a manager."
    )

try:
    selected_llm = create_selected_llm(provider_name, selected_model, api_key_input)
except Exception as exc:
    selected_llm = None
    st.error(f"Could not configure {provider_name}: {exc}")

if selected_llm is None:
    if provider_config["requires_key"]:
        st.warning(f"Add a {provider_name} API key in the sidebar to generate or run a workflow.")
    else:
        st.warning("Start Ollama locally and pull the selected model before generating or running a workflow.")

if "saved_agents" not in st.session_state:
    st.session_state.saved_agents = load_saved_agents()
if "generated_agents" not in st.session_state:
    st.session_state.generated_agents = []
if "generated_tasks" not in st.session_state:
    st.session_state.generated_tasks = []
if "workflow_accepted" not in st.session_state:
    st.session_state.workflow_accepted = False
if "workflow_brief" not in st.session_state:
    st.session_state.workflow_brief = ""
if "history" not in st.session_state:
    st.session_state.history = load_history()

st.divider()
st.header("1. Describe your workflow")
workflow_brief = st.text_area(
    "What should this workflow do?",
    value=st.session_state.workflow_brief,
    height=150,
    placeholder="Example: Build a customer research workflow for a new fintech app. The flow should compare competitors, summarize customer pain points, and recommend feature priorities.",
)
st.session_state.workflow_brief = workflow_brief

generated_text = " ".join(
    [
        agent.get("goal", "") + " " + agent.get("backstory", "")
        for agent in st.session_state.generated_agents
    ]
    + [
        task.get("description", "") + " " + task.get("expected_output", "")
        for task in st.session_state.generated_tasks
    ]
)
web_search_required = workflow_requires_web_search(workflow_brief + " " + generated_text)

if "serper_api_key" not in st.session_state:
    st.session_state.serper_api_key = ""

if web_search_required:
    with st.sidebar:
        st.divider()
        st.subheader("Web research")
        st.caption("This workflow appears to need live website or search access.")
        st.session_state.serper_api_key = st.text_input(
            "Serper API key",
            value=st.session_state.serper_api_key,
            type="password",
            help="Used only during this Streamlit session. It is not saved to a file.",
        )

with st.expander("Workflow history log", expanded=False):
    if st.session_state.history:
        for index, item in enumerate(st.session_state.history):
            prompt_preview = item["prompt"]
            label = f"{item.get('created_at', 'Unknown')} - {prompt_preview[:80]}"
            if st.button(label, key=f"history_{index}"):
                st.session_state.workflow_brief = item["prompt"]
                st.session_state.generated_agents = []
                st.session_state.generated_tasks = []
                st.session_state.workflow_accepted = False
                st.rerun()
    else:
        st.caption("No saved workflow history yet.")

if st.button("Generate workflow", type="primary"):
    if not workflow_brief.strip():
        st.error("Please describe the workflow you want to build.")
    elif selected_llm is None:
        st.error(f"Configure {provider_name} before generating the workflow.")
    else:
        try:
            with st.spinner("Generating agents and tasks from your brief..."):
                generated_agents, generated_tasks = generate_workflow_from_brief(workflow_brief, selected_llm)
                st.session_state.generated_agents = generated_agents
                st.session_state.generated_tasks = generated_tasks
                st.session_state.workflow_accepted = False
                add_history_entry(workflow_brief)
            st.success("A draft workflow was generated.")
        except Exception as exc:
            st.error(f"Could not generate the workflow: {exc}")

if st.session_state.generated_agents or st.session_state.generated_tasks:
    st.divider()
    st.header("2. Review generated plan")

    st.subheader("Suggested agents")
    for agent in st.session_state.generated_agents:
        st.write(f"**{agent['role']}**")
        st.caption(agent["goal"])
        st.caption(agent["backstory"])

    st.subheader("Suggested tasks")
    for index, task in enumerate(st.session_state.generated_tasks, start=1):
        st.write(f"**Task {index}:** {task['agent_role']}")
        st.write(task["description"])
        st.caption(f"Expected output: {task['expected_output']}")

    if st.button("Use this generated workflow"):
        st.session_state.saved_agents = st.session_state.saved_agents + [
            agent for agent in st.session_state.generated_agents
            if agent not in st.session_state.saved_agents
        ]
        save_agents(st.session_state.saved_agents)
        st.session_state.workflow_accepted = True
        st.success("The generated workflow is ready to run.")

st.divider()
st.header("3. Run workflow")

uploaded_document = st.file_uploader(
    "Attach a source document for the workflow",
    type=["txt", "md", "csv", "json", "pdf", "docx"],
    help="Optional context to include in every task.",
)

document_text = ""
if uploaded_document:
    try:
        document_text = extract_document_text(uploaded_document)
        st.success(f"Attached {uploaded_document.name} ({len(document_text):,} characters).")
        with st.expander("Preview attached document"):
            st.text(document_text[:5000] or "No readable text was found.")
    except Exception as exc:
        st.error(f"Could not read {uploaded_document.name}: {exc}")

active_agents = st.session_state.generated_agents if st.session_state.workflow_accepted else []
active_tasks = st.session_state.generated_tasks if st.session_state.workflow_accepted else []

if not active_agents or not active_tasks:
    st.info("Generate a workflow first and accept it to start the run.")
else:
    if st.button("Run workflow", type="primary"):
        if not selected_llm:
            st.error(f"Configure {provider_name} before running.")
        elif web_search_required and not st.session_state.serper_api_key.strip():
            st.error("This workflow needs web research. Add the Serper API key in the left sidebar before running.")
        else:
            previous_serper_key = os.environ.get("SERPER_API_KEY")
            try:
                web_search_tool = None
                if web_search_required:
                    from crewai_tools import SerperDevTool

                    os.environ["SERPER_API_KEY"] = st.session_state.serper_api_key.strip()
                    web_search_tool = SerperDevTool()

                with st.spinner("Running your generated workflow..."):
                    created_agents = {}
                    crewai_agents = []
                    for ag in active_agents:
                        agent_obj = Agent(
                            role=ag["role"],
                            goal=ag["goal"],
                            backstory=ag["backstory"],
                            llm=selected_llm,
                            tools=[web_search_tool] if web_search_tool else [],
                            verbose=False,
                        )
                        created_agents[ag["role"]] = agent_obj
                        crewai_agents.append(agent_obj)

                    crewai_tasks = []
                    for task in active_tasks:
                        assigned_agent = created_agents.get(task["agent_role"])
                        if assigned_agent is None:
                            assigned_agent = next(iter(created_agents.values()))

                        task_description = task["description"]
                        if document_text:
                            task_description = (
                                f"{task_description}\n\n"
                                "SOURCE DOCUMENT (Use this as reference material):\n"
                                f"{document_text}"
                            )

                        task_obj = Task(
                            description=task_description,
                            expected_output=task["expected_output"],
                            agent=assigned_agent,
                        )
                        crewai_tasks.append(task_obj)

                    selected_process = Process.sequential if process_type == "Sequential" else Process.hierarchical
                    crew = Crew(
                        agents=crewai_agents,
                        tasks=crewai_tasks,
                        process=selected_process,
                        manager_llm=selected_llm if selected_process == Process.hierarchical else None,
                        verbose=False,
                    )
                    result = crew.kickoff()

                st.success("Workflow completed.")
                st.markdown("### Final result")
                st.markdown(result.raw)
                st.download_button(
                    "Download result as Markdown",
                    data=result.raw,
                    file_name="workflow_result.md",
                    mime="text/markdown",
                )
                try:
                    st.download_button(
                        "Download result as Word document",
                        data=create_docx_download(result.raw),
                        file_name="workflow_result.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                except ImportError:
                    st.caption("Install python-docx to enable Word document downloads.")
            except Exception:
                st.error("The workflow could not be completed. Check your settings and try again.")
            finally:
                if previous_serper_key is None:
                    os.environ.pop("SERPER_API_KEY", None)
                else:
                    os.environ["SERPER_API_KEY"] = previous_serper_key
