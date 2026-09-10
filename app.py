import json
import os
from io import BytesIO
from pathlib import Path

import streamlit as st
from crewai import Agent, Task, Crew, Process, LLM

AGENTS_FILE = Path(__file__).with_name("agents.json")
DEFAULT_AGENTS = [
    {"role": "Research Analyst", "goal": "Find reliable information and turn it into clear, useful insights", "backstory": "You are a careful researcher who compares sources, spots patterns, and explains complex topics plainly."},
    {"role": "Content Strategist", "goal": "Shape information into focused content for a specific audience", "backstory": "You are an experienced editor who understands audience needs, structure, tone, and persuasive communication."},
    {"role": "Data Analyst", "goal": "Interpret data, identify meaningful trends, and communicate practical conclusions", "backstory": "You are a methodical analyst who checks assumptions, looks for outliers, and presents evidence-driven recommendations."},
    {"role": "Quality Reviewer", "goal": "Check work for accuracy, completeness, clarity, and actionable improvements", "backstory": "You are a constructive reviewer who catches gaps and inconsistencies while keeping the final result easy to use."},
]


def load_saved_agents():
    if not AGENTS_FILE.exists():
        return [dict(agent) for agent in DEFAULT_AGENTS]
    try:
        with AGENTS_FILE.open("r", encoding="utf-8") as file:
            saved_agents = json.load(file)
        return saved_agents if isinstance(saved_agents, list) else []
    except (OSError, json.JSONDecodeError):
        return [dict(agent) for agent in DEFAULT_AGENTS]


def save_agents(agents):
    with AGENTS_FILE.open("w", encoding="utf-8") as file:
        json.dump(agents, file, indent=2)


def extract_document_text(uploaded_file):
    """Extract text from the supported document formats."""
    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith((".txt", ".md", ".csv", ".json")):
        return file_bytes.decode("utf-8", errors="replace")

    if file_name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if file_name.endswith(".docx"):
        from docx import Document

        document = Document(BytesIO(file_bytes))
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs).strip()

    raise ValueError("Unsupported file type. Use TXT, MD, CSV, JSON, PDF, or DOCX.")


def create_docx_download(text):
    from docx import Document

    document = Document()
    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


st.set_page_config(page_title="Workflow Studio", layout="wide")
st.title("Workflow Studio")
st.subheader("Build and run focused workflows with reusable specialist roles")

# Sidebar: API Configuration & Controls
with st.sidebar:
    st.header("Model settings")
    api_key_input = st.text_input(
        "Provider API key",
        type="password", 
        value=os.getenv("GEMINI_API_KEY", ""),
        help="Your provider API key is used only when a workflow runs."
    )
    
    gemini_model = st.selectbox(
        "Select model",
        options=["gemini/gemini-3.5-flash", "gemini/gemini-3.5-pro", "gemini/gemini-3.5-flash-lite"],
        index=0
    )
    
    process_type = st.radio(
        "Workflow mode",
        options=["Sequential", "Hierarchical"],
        index=0,
        help="Choose whether tasks run in order or are coordinated by a manager."
    )

if api_key_input:
    gemini_llm = LLM(
        model=gemini_model,
        api_key=api_key_input
    )
else:
    gemini_llm = None
    st.warning("Add a provider API key in the sidebar to run a workflow.")

if "saved_agents" not in st.session_state:
    st.session_state.saved_agents = load_saved_agents()
if "agents" not in st.session_state:
    st.session_state.agents = []
if "tasks" not in st.session_state:
    st.session_state.tasks = []
if "editing_agent_index" not in st.session_state:
    st.session_state.editing_agent_index = None
if "editing_task_index" not in st.session_state:
    st.session_state.editing_task_index = None

st.divider()
st.header("1. Build your role library")
st.caption("Save specialist roles once and reuse them in future sessions.")

col1, col2 = st.columns([2, 1])

with col1:
    saved_roles = [agent["role"] for agent in st.session_state.saved_agents]
    if saved_roles:
        selected_role = st.selectbox("Choose a saved role", saved_roles)
        if st.button("Load role into workflow"):
            selected_agent = next(agent for agent in st.session_state.saved_agents if agent["role"] == selected_role)
            if not any(agent["role"] == selected_role for agent in st.session_state.agents):
                st.session_state.agents.append(dict(selected_agent))
                st.success(f"'{selected_role}' is ready to use.")
            else:
                st.info(f"'{selected_role}' is already in this workflow.")

    with st.form("add_agent_form", clear_on_submit=True):
        role = st.text_input("Role", placeholder="e.g., Product Researcher")
        goal = st.text_input("Goal", placeholder="e.g., Compare options and recommend the best fit")
        backstory = st.text_area("Working style", placeholder="e.g., You are a thoughtful specialist who uses evidence and explains tradeoffs.")
        
        submitted_agent = st.form_submit_button("Save role")
        
        if submitted_agent:
            if role and goal and backstory:
                agent_dict = {
                    "role": role,
                    "goal": goal,
                    "backstory": backstory
                }
                st.session_state.saved_agents = [
                    agent for agent in st.session_state.saved_agents if agent["role"] != role
                ] + [agent_dict]
                save_agents(st.session_state.saved_agents)
                st.session_state.agents.append(dict(agent_dict))
                st.success(f"'{role}' was saved to your role library.")
            else:
                st.error("Please fill in the role, goal, and working style.")

with col2:
    st.subheader("Workflow roles")
    if not st.session_state.agents:
        st.info("Load a saved role to get started.")
    else:
        for index, agent in enumerate(st.session_state.agents):
            st.write(f"**{agent['role']}**")
            col_edit, col_delete = st.columns([1, 1])
            with col_edit:
                if st.button("Edit", key=f"edit_agent_{index}"):
                    st.session_state.editing_agent_index = index
                    st.rerun()
            with col_delete:
                if st.button("Delete", key=f"delete_agent_{index}"):
                    removed_role = agent["role"]
                    st.session_state.agents.pop(index)
                    st.session_state.saved_agents = [
                        saved_agent for saved_agent in st.session_state.saved_agents
                        if saved_agent["role"] != removed_role
                    ]
                    save_agents(st.session_state.saved_agents)
                    st.session_state.tasks = [
                        task for task in st.session_state.tasks if task["agent_role"] != removed_role
                    ]
                    if st.session_state.editing_agent_index == index:
                        st.session_state.editing_agent_index = None
                    for task_index, task in enumerate(st.session_state.tasks):
                        task["id"] = task_index + 1
                    st.success(f"'{removed_role}' was removed from the workflow.")
                    st.rerun()

        if st.session_state.editing_agent_index is not None:
            agent_index = st.session_state.editing_agent_index
            current_agent = st.session_state.agents[agent_index]
            with st.form("edit_agent_form", clear_on_submit=True):
                edited_role = st.text_input("Role", value=current_agent["role"])
                edited_goal = st.text_input("Goal", value=current_agent["goal"])
                edited_backstory = st.text_area("Working style", value=current_agent["backstory"])

                col_save, col_cancel = st.columns(2)
                with col_save:
                    save_agent = st.form_submit_button("Save changes")
                with col_cancel:
                    cancel_agent = st.form_submit_button("Cancel")

                if save_agent:
                    if edited_role and edited_goal and edited_backstory:
                        previous_role = current_agent["role"]
                        updated_agent = {
                            "role": edited_role,
                            "goal": edited_goal,
                            "backstory": edited_backstory,
                        }
                        st.session_state.agents[agent_index] = dict(updated_agent)
                        st.session_state.saved_agents = [
                            updated_agent if saved_agent.get("role") == previous_role else saved_agent
                            for saved_agent in st.session_state.saved_agents
                        ]
                        for task in st.session_state.tasks:
                            if task["agent_role"] == previous_role:
                                task["agent_role"] = edited_role
                        save_agents(st.session_state.saved_agents)
                        st.session_state.editing_agent_index = None
                        st.success(f"'{previous_role}' was updated.")
                        st.rerun()
                    else:
                        st.error("Please fill in the role, goal, and working style.")
                elif cancel_agent:
                    st.session_state.editing_agent_index = None
                    st.rerun()

    if st.button("Restore default roles"):
        existing_roles = {agent["role"] for agent in st.session_state.saved_agents}
        missing_defaults = [
            dict(agent) for agent in DEFAULT_AGENTS if agent["role"] not in existing_roles
        ]
        st.session_state.saved_agents = st.session_state.saved_agents + missing_defaults
        save_agents(st.session_state.saved_agents)
        st.success("Default roles restored. Your custom roles were kept.")
        st.rerun()
    if saved_roles and st.button("Delete selected saved role"):
        st.session_state.saved_agents = [
            agent for agent in st.session_state.saved_agents if agent["role"] != selected_role
        ]
        save_agents(st.session_state.saved_agents)
        st.rerun()

st.divider()
st.header("2. Define your tasks")

uploaded_document = st.file_uploader(
    "Attach a source document for the agents",
    type=["txt", "md", "csv", "json", "pdf", "docx"],
    help="The document text is provided to every task as shared context.",
)
document_text = ""
if uploaded_document:
    try:
        document_text = extract_document_text(uploaded_document)
        st.success(f"Attached {uploaded_document.name} ({len(document_text):,} characters).")
        with st.expander("Preview attached document"):
            st.text(document_text[:5000] or "No readable text was found.")
    except Exception as exc:
        st.error(f"Could not read {uploaded_document.name}: {exc}")

if not st.session_state.agents:
    st.info("Load at least one role above before creating tasks.")
else:
    col_t1, col_t2 = st.columns([2, 1])
    
    agent_roles = [ag["role"] for ag in st.session_state.agents]

    with col_t1:
        with st.form("add_task_form", clear_on_submit=True):
            task_desc = st.text_area("Task Description", placeholder="e.g., Analyze recent trends in multi-agent orchestration platforms.")
            expected_output = st.text_input("Expected Output", placeholder="e.g., A 3-bullet-point summary report.")
            assigned_agent_role = st.selectbox("Assign to Agent", options=agent_roles)
            
            submitted_task = st.form_submit_button("Add task")
            
            if submitted_task:
                if task_desc and expected_output:
                    task_dict = {
                        "id": len(st.session_state.tasks) + 1,
                        "description": task_desc,
                        "expected_output": expected_output,
                        "agent_role": assigned_agent_role
                    }
                    st.session_state.tasks.append(task_dict)
                    st.success(f"Task #{task_dict['id']} added.")
                else:
                    st.error("Please fill in the task description and expected output.")

    with col_t2:
        st.subheader("Current tasks")
        if not st.session_state.tasks:
            st.info("No tasks created yet.")
        else:
            for index, task in enumerate(st.session_state.tasks):
                st.write(f"**Task {task['id']}:** assigned to `{task['agent_role']}`")
                col_edit_task, col_delete_task = st.columns([1, 1])
                with col_edit_task:
                    if st.button("Edit task", key=f"edit_task_{task['id']}"):
                        st.session_state.editing_task_index = index
                        st.rerun()
                with col_delete_task:
                    if st.button("Delete task", key=f"delete_task_{task['id']}"):
                        removed_task_id = task["id"]
                        st.session_state.tasks.pop(index)
                        for updated_task_index, updated_task in enumerate(st.session_state.tasks):
                            updated_task["id"] = updated_task_index + 1
                        if st.session_state.editing_task_index == index:
                            st.session_state.editing_task_index = None
                        st.success(f"Task #{removed_task_id} was deleted.")
                        st.rerun()

            if st.session_state.editing_task_index is not None:
                task_index = st.session_state.editing_task_index
                current_task = st.session_state.tasks[task_index]
                with st.form("edit_task_form", clear_on_submit=True):
                    edited_description = st.text_area("Task Description", value=current_task["description"])
                    edited_output = st.text_input("Expected Output", value=current_task["expected_output"])
                    edited_agent_role = st.selectbox(
                        "Assign to Agent",
                        options=[ag["role"] for ag in st.session_state.agents],
                        index=[ag["role"] for ag in st.session_state.agents].index(current_task["agent_role"]),
                    )

                    col_save_task, col_cancel_task = st.columns(2)
                    with col_save_task:
                        save_task = st.form_submit_button("Save task")
                    with col_cancel_task:
                        cancel_task = st.form_submit_button("Cancel")

                    if save_task:
                        if edited_description and edited_output:
                            st.session_state.tasks[task_index]["description"] = edited_description
                            st.session_state.tasks[task_index]["expected_output"] = edited_output
                            st.session_state.tasks[task_index]["agent_role"] = edited_agent_role
                            st.session_state.editing_task_index = None
                            st.success(f"Task #{current_task['id']} updated.")
                            st.rerun()
                        else:
                            st.error("Please fill in the task description and expected output.")
                    elif cancel_task:
                        st.session_state.editing_task_index = None
                        st.rerun()

st.divider()
st.header("3. Run your workflow")

if st.button("Run workflow", type="primary"):
    if not gemini_llm:
        st.error("Add a provider API key before running.")
    elif not st.session_state.agents:
        st.error("Load at least one saved role.")
    elif not st.session_state.tasks:
        st.error("Please add at least one task.")
    else:
        try:
            with st.spinner("Running your workflow..."):
                # Step A: Instantiate CrewAI Agent objects dynamically
                created_agents = {}
                crewai_agents = []
                
                for ag in st.session_state.agents:
                    agent_obj = Agent(
                        role=ag["role"],
                        goal=ag["goal"],
                        backstory=ag["backstory"],
                        llm=gemini_llm,  # Attach Gemini LLM
                        verbose=False
                    )
                    created_agents[ag["role"]] = agent_obj
                    crewai_agents.append(agent_obj)

                # Step B: Instantiate CrewAI Task objects dynamically
                crewai_tasks = []
                for t in st.session_state.tasks:
                    assigned_agent = created_agents[t["agent_role"]]
                    task_description = t["description"]
                    if document_text:
                        task_description = (
                            f"{task_description}\n\n"
                            "SOURCE DOCUMENT (use this as the material to read and edit):\n"
                            f"{document_text}"
                        )
                    task_obj = Task(
                        description=task_description,
                        expected_output=t["expected_output"],
                        agent=assigned_agent
                    )
                    crewai_tasks.append(task_obj)

                # Step C: Select Execution Process
                selected_process = Process.sequential if process_type == "Sequential" else Process.hierarchical

                # Step D: Assemble Crew
                crew = Crew(
                    agents=crewai_agents,
                    tasks=crewai_tasks,
                    process=selected_process,
                    manager_llm=gemini_llm if selected_process == Process.hierarchical else None,
                    verbose=False
                )

                # Step E: Kickoff Workflow
                result = crew.kickoff()

            st.success("Workflow completed.")
            st.markdown("### Final result")
            st.markdown(result.raw)
            st.download_button(
                "Download result as Markdown",
                data=result.raw,
                file_name="workflow_result.md",
                mime="text/markdown",
            )
            try:
                st.download_button(
                    "Download result as Word document",
                    data=create_docx_download(result.raw),
                    file_name="workflow_result.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            except ImportError:
                st.caption("Install python-docx to enable Word document downloads.")

        except Exception:
            st.error("The workflow could not be completed. Check your settings and try again.")